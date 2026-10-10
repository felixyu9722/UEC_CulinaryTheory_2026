import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('unit20_quiz .docx')
print(f"Total paragraphs: {len(doc.paragraphs)}")
for idx, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"[{idx}] {p.text}")

print(f"\nTotal tables: {len(doc.tables)}")
for t_idx, tbl in enumerate(doc.tables):
    print(f"--- Table {t_idx} ({len(tbl.rows)} rows, {len(tbl.columns)} cols) ---")
    for r in tbl.rows:
        print([c.text.strip() for c in r.cells])
