# -*- coding: utf-8 -*-
"""
generate_quiz_2026_b.py
Complete production generator for 2026 UEC Senior 1 Culinary Arts (Baking Theory) Test Bank (Version B - B版).
Style: Authentic UEC Past Exam Style (统考历届真题风格), directly matching unit20_quiz .docx.
Features:
- Natural question & option phrasing (NO sentence length constraints whatsoever!)
- Question format: 1．（　　）题干
- Options: Tab-separated when short, clean lines when long
- Student paper: Exam header, student info, instructions, 15 MCQs
- Teacher paper: Exam header, compact Answer & Explanation Table (题号 | 答案 | 解析), followed by detailed question review
- 20 Single-unit Student Papers (.docx)
- 20 Single-unit Teacher Papers (.docx)
- 1 Master 300-Question Student Compilation (.docx)
- 1 Master 300-Question Teacher Compilation (.docx)
- 1 Standard JSON dataset (.json)
- 1 Comprehensive Markdown Guide (.md)
Target folder: 厨艺理论测验题2026 (Preserves all existing Version A files)
"""

import os, sys, json, re
sys.stdout.reconfigure(encoding='utf-8')

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls

from data_b_part1 import B_UNITS_PART1
from data_b_part2 import B_UNITS_PART2
from data_b_part3 import B_UNITS_PART3
from data_b_part4 import B_UNITS_PART4

WORKSPACE_ROOT = r'c:\Users\GOH\Desktop\2026年统考厨艺理论'
OUT_DIR = os.path.join(WORKSPACE_ROOT, '厨艺理论测验题2026')
os.makedirs(OUT_DIR, exist_ok=True)

ALL_B_UNITS = B_UNITS_PART1 + B_UNITS_PART2 + B_UNITS_PART3 + B_UNITS_PART4

# -------------------------------------------------------------
# XML STYLING UTILITIES
# -------------------------------------------------------------
def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def create_b_document(header_text=""):
    doc = Document()
    for s in doc.sections:
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.0)
        s.right_margin = Cm(2.0)
        
        if header_text:
            hdr = s.header.paragraphs[0]
            hdr.text = header_text
            hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if hdr.runs:
                hdr.runs[0].font.name = "Microsoft YaHei"
                hdr.runs[0].font.size = Pt(8.5)
                hdr.runs[0].font.color.rgb = RGBColor(148, 163, 184)

    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor(30, 41, 59)
    style.paragraph_format.line_spacing = 1.25
    style.paragraph_format.space_after = Pt(2)
    rPr = style.element.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="Microsoft YaHei" w:ascii="Arial" w:hAnsi="Arial"/>')
    rPr.append(rFonts)

    return doc

def format_b_options(doc, opts):
    lens = [len(opts[k]) for k in ['A', 'B', 'C', 'D']]
    max_len = max(lens)

    if max_len <= 15:
        # 2 lines with tab separation
        p1 = doc.add_paragraph()
        p1.paragraph_format.left_indent = Inches(0.25)
        p1.paragraph_format.space_after = Pt(1.5)
        p1.add_run(f"A．{opts['A']}\tB．{opts['B']}")

        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Inches(0.25)
        p2.paragraph_format.space_after = Pt(3.5)
        p2.add_run(f"C．{opts['C']}\tD．{opts['D']}")
    else:
        # 4 separate lines
        for k in ['A', 'B', 'C', 'D']:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(1.5)
            p.add_run(f"{k}．{opts[k]}")
        # Add slight space after question block
        p.paragraph_format.space_after = Pt(3.5)

# -------------------------------------------------------------
# 1. GENERATE STUDENT UNIT DOC (B版)
# -------------------------------------------------------------
def build_student_b_doc(unit, doc=None, is_standalone=True):
    if doc is None:
        hdr = f"2026年马来西亚华文独中高中统考（UEC）厨艺科 · 随堂测验卷（B版）"
        doc = create_b_document(header_text=hdr)

    # Unit Title
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(4)
    p_t.paragraph_format.space_after = Pt(2)
    r_t = p_t.add_run(f"第{unit['unit_num']}单元　{unit['title']}")
    r_t.font.size = Pt(18)
    r_t.font.bold = True
    r_t.font.name = "Arial"

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(8)
    r_sub = p_sub.add_run("单元测验卷（B版 · 统考真题风格）")
    r_sub.font.size = Pt(14)
    r_sub.font.bold = True
    r_sub.font.name = "Arial"

    # Student Info Line
    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_after = Pt(3)
    p_info.add_run("姓名：______________　班级：__________　日期：__________　得分：________ / 30")

    # Time & Meta
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(6)
    p_meta.add_run("时间：25 分钟　　题型：单项选择题 15 题（每题 2 分，共 30 分）")

    # Instructions Box
    p_inst_h = doc.add_paragraph()
    p_inst_h.paragraph_format.space_before = Pt(4)
    p_inst_h.paragraph_format.space_after = Pt(2)
    r_ih = p_inst_h.add_run("注意事项")
    r_ih.font.bold = True

    p_inst_c = doc.add_paragraph()
    p_inst_c.paragraph_format.space_after = Pt(8)
    p_inst_c.add_run("每题只有一个正确答案，请把答案字母填在题号后的括号内。")

    # Section Title
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(6)
    p_sec.paragraph_format.space_after = Pt(6)
    r_sec = p_sec.add_run("第一部分　选择题（每题 2 分，共 30 分）")
    r_sec.font.size = Pt(13)
    r_sec.font.bold = True
    r_sec.font.name = "Arial"

    # Questions
    for q in unit['questions']:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_after = Pt(2)
        # 1．（　　）题干
        p_q.add_run(f"{q['num']}．（　　）{q['q']}")
        format_b_options(doc, q['opts'])

    if is_standalone:
        out_path = os.path.join(OUT_DIR, f"第{unit['unit_num']:02d}单元_{unit['title']}_学生测验卷_B版.docx")
        doc.save(out_path)
        print(f"Generated Student Doc (B-Ver): {out_path}")

    return doc

# -------------------------------------------------------------
# 2. GENERATE TEACHER UNIT DOC (B版)
# -------------------------------------------------------------
def build_teacher_b_doc(unit, doc=None, is_standalone=True):
    if doc is None:
        hdr = f"2026年马来西亚华文独中高中统考（UEC）厨艺科 · 教师参考答案（B版）"
        doc = create_b_document(header_text=hdr)

    # Unit Title
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(4)
    p_t.paragraph_format.space_after = Pt(2)
    r_t = p_t.add_run(f"第{unit['unit_num']}单元　{unit['title']}")
    r_t.font.size = Pt(18)
    r_t.font.bold = True
    r_t.font.name = "Arial"

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(8)
    r_sub = p_sub.add_run("单元测验卷 · 教师参考答案与解析（B版 · 统考真题风格）")
    r_sub.font.size = Pt(14)
    r_sub.font.bold = True
    r_sub.font.name = "Arial"

    # Meta
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(10)
    p_meta.add_run("【教师专用参考】 满分：30 分（共 15 题，每题 2 分）")

    # Table: 题号 | 答案 | 解析 (Identical style to unit20_quiz .docx Table 6!)
    tbl = doc.add_table(rows=16, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Table Header
    headers = ["题号", "答案", "解析"]
    for c_i, h in enumerate(headers):
        c = tbl.rows[0].cells[c_i]
        c.text = h
        set_cell_background(c, "F1F5F9")
        set_cell_margins(c, top=80, bottom=80, left=80, right=80)
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.size = Pt(9.5)
        if c_i in [0, 1]:
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Table Rows
    for q_idx, q in enumerate(unit['questions']):
        r = tbl.rows[q_idx + 1]
        
        # Col 0: 题号
        r.cells[0].text = str(q['num'])
        set_cell_margins(r.cells[0], top=60, bottom=60, left=60, right=60)
        r.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r.cells[0].paragraphs[0].runs[0].font.size = Pt(9.5)

        # Col 1: 答案
        r.cells[1].text = q['ans']
        set_cell_margins(r.cells[1], top=60, bottom=60, left=60, right=60)
        r.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r.cells[1].paragraphs[0].runs[0].font.size = Pt(10)
        r.cells[1].paragraphs[0].runs[0].font.bold = True
        r.cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(21, 128, 61)

        # Col 2: 解析
        r.cells[2].text = q['exp']
        set_cell_margins(r.cells[2], top=60, bottom=60, left=80, right=80)
        r.cells[2].paragraphs[0].runs[0].font.size = Pt(9)
        r.cells[2].paragraphs[0].runs[0].font.color.rgb = RGBColor(51, 65, 85)

    set_table_borders(tbl, color="CBD5E1")
    doc.add_paragraph() # Spacer

    # Section 2: Detailed Questions Review
    p_dt = doc.add_paragraph()
    p_dt.paragraph_format.space_before = Pt(8)
    p_dt.paragraph_format.space_after = Pt(6)
    r_dt = p_dt.add_run("第二部分　全题题干与选项答案核对")
    r_dt.font.size = Pt(12)
    r_dt.font.bold = True
    r_dt.font.name = "Arial"

    for q in unit['questions']:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_after = Pt(2)
        r_num = p_q.add_run(f"{q['num']}．（ ")
        r_ans = p_q.add_run(f" {q['ans']} ")
        r_ans.font.bold = True
        r_ans.font.color.rgb = RGBColor(220, 38, 38)
        r_rest = p_q.add_run(f" ）{q['q']}")
        format_b_options(doc, q['opts'])

    if is_standalone:
        out_path = os.path.join(OUT_DIR, f"第{unit['unit_num']:02d}单元_{unit['title']}_教师解析与参考答案_B版.docx")
        doc.save(out_path)
        print(f"Generated Teacher Doc (B-Ver): {out_path}")

    return doc

# -------------------------------------------------------------
# 3. GENERATE MASTER COMPILATIONS (B版 300 题)
# -------------------------------------------------------------
def build_master_student_b(all_units):
    hdr = "2026年马来西亚华文独中高中统考（UEC）厨艺理论 · 全套20单元300题测验题总汇编（B版）"
    doc = create_b_document(header_text=hdr)

    # Master Cover
    p_cov = doc.add_paragraph()
    p_cov.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov.paragraph_format.space_before = Pt(72)
    p_cov.paragraph_format.space_after = Pt(12)

    r0 = p_cov.add_run("2026年马来西亚华文独立中学高中统考（UEC）\n")
    r0.font.size = Pt(15)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(100, 116, 139)

    r1 = p_cov.add_run("高中厨艺科 · 烘焙理论\n高一全套第 01 至 20 单元测验题卷总汇编（B版）\n")
    r1.font.size = Pt(22)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(30, 58, 138)

    r2 = p_cov.add_run("（300 题统考真题风格 · 随堂测验标准卷）\n\n")
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(180, 83, 9)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(48)
    p_meta.add_run(
        "【版本说明】：B版（统考历届真题风格，题干与选项自然流畅，无句子长度约束限制）\n"
        "【试卷架构】：全套 20 个完整单元 ｜ 每单元 15 题选择题 ｜ 总计 300 题 ｜ 满分 600 分\n"
        "【编写参考】：2020～2025年统考历届真题、Notion 2026 高一讲义、董总课纲"
    )
    doc.add_page_break()

    # Table of Contents
    h_toc = doc.add_paragraph()
    h_toc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_th = h_toc.add_run("高一《烘焙理论》全套 20 单元目录索引（B版）")
    r_th.font.size = Pt(16)
    r_th.font.bold = True

    tbl_toc = doc.add_table(rows=21, cols=4)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    toc_h = ["单元序号", "课程单元主题", "题目数量", "建议时量"]
    for c_i, h in enumerate(toc_h):
        c = tbl_toc.rows[0].cells[c_i]
        c.text = h
        set_cell_background(c, "F1F5F9")
        set_cell_margins(c, top=80, bottom=80, left=80, right=80)
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for idx, u in enumerate(all_units):
        row = tbl_toc.rows[idx+1]
        row.cells[0].text = f"第 {u['unit_num']:02d} 单元"
        row.cells[1].text = u['title']
        row.cells[2].text = "15 题选择题"
        row.cells[3].text = "25 分钟"
        for c_i in range(4):
            set_cell_margins(row.cells[c_i], top=60, bottom=60, left=80, right=80)
            row.cells[c_i].paragraphs[0].runs[0].font.size = Pt(9)
            if c_i in [0, 2, 3]:
                row.cells[c_i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    set_table_borders(tbl_toc, color="CBD5E1")
    doc.add_page_break()

    # Append units
    for idx, u in enumerate(all_units):
        build_student_b_doc(u, doc=doc, is_standalone=False)
        if idx < len(all_units) - 1:
            doc.add_page_break()

    out_path = os.path.join(OUT_DIR, "高一厨艺理论_全套第01至20单元_300题学生测验卷总汇编_B版.docx")
    doc.save(out_path)
    print(f"Generated Master Student Compilation (B-Ver): {out_path}")

def build_master_teacher_b(all_units):
    hdr = "2026年马来西亚华文独中高中统考（UEC）厨艺理论 · 教师专用解析与考点总汇编（B版）"
    doc = create_b_document(header_text=hdr)

    # Master Cover
    p_cov = doc.add_paragraph()
    p_cov.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov.paragraph_format.space_before = Pt(72)
    p_cov.paragraph_format.space_after = Pt(12)

    r0 = p_cov.add_run("2026年马来西亚华文独立中学高中统考（UEC）\n")
    r0.font.size = Pt(15)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(100, 116, 139)

    r1 = p_cov.add_run("高中厨艺科 · 烘焙理论\n高一全套第 01 至 20 单元教师解析总汇编（B版）\n")
    r1.font.size = Pt(22)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(21, 128, 61)

    r2 = p_cov.add_run("（300 题全真答案速查矩阵 ＆ 逐题考点解析手册）\n\n")
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(180, 83, 9)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(48)
    p_meta.add_run(
        "【配套卷别】：高一烘焙理论随堂测验卷（B版 · 统考真题风格 · 300 题）\n"
        "【特色功能】：含 20 单元 300 题答案速查总矩阵 ＋ 逐单元紧凑答题解析表 ＋ 全题题干核对\n"
        "【编写团队】：2026 年马来西亚华文独中高中统考厨艺科教研组"
    )
    doc.add_page_break()

    # Master Quick Answer Matrix (20 units x 15 questions)
    h_m = doc.add_paragraph()
    h_m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_mh = h_m.add_run("全套 20 单元 300 题答案速查总矩阵 (Master Key Matrix · B版)")
    r_mh.font.size = Pt(15)
    r_mh.font.bold = True
    r_mh.font.color.rgb = RGBColor(30, 58, 138)

    tbl_matrix = doc.add_table(rows=21, cols=16)
    tbl_matrix.alignment = WD_TABLE_ALIGNMENT.CENTER

    tbl_matrix.rows[0].cells[0].text = "单元"
    set_cell_background(tbl_matrix.rows[0].cells[0], "F1F5F9")
    set_cell_margins(tbl_matrix.rows[0].cells[0], top=60, bottom=60, left=40, right=40)
    tbl_matrix.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tbl_matrix.rows[0].cells[0].paragraphs[0].runs[0].font.size = Pt(8.5)
    tbl_matrix.rows[0].cells[0].paragraphs[0].runs[0].font.bold = True

    for i in range(1, 16):
        c = tbl_matrix.rows[0].cells[i]
        c.text = f"Q{i}"
        set_cell_background(c, "F1F5F9")
        set_cell_margins(c, top=60, bottom=60, left=40, right=40)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].runs[0].font.size = Pt(8.5)
        c.paragraphs[0].runs[0].font.bold = True

    for u_idx, u in enumerate(all_units):
        row = tbl_matrix.rows[u_idx+1]
        row.cells[0].text = f"U{u['unit_num']:02d}"
        set_cell_background(row.cells[0], "F8FAFC")
        set_cell_margins(row.cells[0], top=50, bottom=50, left=40, right=40)
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(8)
        row.cells[0].paragraphs[0].runs[0].font.bold = True

        for q_i in range(1, 16):
            q = u['questions'][q_i-1]
            c = row.cells[q_i]
            c.text = q['ans']
            set_cell_margins(c, top=50, bottom=50, left=40, right=40)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.paragraphs[0].runs[0].font.size = Pt(8.5)
            c.paragraphs[0].runs[0].font.bold = True
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(21, 128, 61)

    set_table_borders(tbl_matrix, color="CBD5E1")
    doc.add_page_break()

    # Append units
    for idx, u in enumerate(all_units):
        build_teacher_b_doc(u, doc=doc, is_standalone=False)
        if idx < len(all_units) - 1:
            doc.add_page_break()

    out_path = os.path.join(OUT_DIR, "高一厨艺理论_全套第01至20单元_300题教师详细解析总汇编_B版.docx")
    doc.save(out_path)
    print(f"Generated Master Teacher Compilation (B-Ver): {out_path}")

# -------------------------------------------------------------
# 4. EXPORT JSON AND MARKDOWN GUIDE (B版)
# -------------------------------------------------------------
def export_b_json_and_md(all_units):
    # 1. JSON
    json_path = os.path.join(OUT_DIR, "高一厨艺理论_20单元300题标准测验题库_B版.json")
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(all_units, f, ensure_ascii=False, indent=2)
    print(f"Exported JSON (B-Ver): {json_path}")

    # 2. Markdown Guide
    md_path = os.path.join(OUT_DIR, "00_高一厨艺理论测验题2026_B版_统考真题风格说明.md")
    lines = [
        "# 2026年马来西亚华文独中高中统考（UEC）厨艺科（烘焙理论）",
        "## 高一全套第01至20单元 · 300题测验题库【B版 · 统考真题风格】架构说明\n",
        "> **核心定位**：完全参考真实独中统考真题排版与题风，贴合后厨实际教学，取消一切选项句子长度约束。  ",
        "> **参考样本**：直接融入老师提供的 `unit20_quiz .docx` 范本文档格式与经典测验风格。  ",
        "> **成果范围**：全套 20 个单元，每单元 15 题选择题，共 300 题。含 20 份独立学生卷、20 份独立教师解析卷、2 份 300 题全套总汇编大卷。\n",
        "---",
        "### 一、B版命题原则与特色（对比A版）\n",
        "1. **完全废除“选项句子长度约束”**：",
        "   - 彻底不再使用字数补齐或刻意对称约束；选项完全依据真实考题，长短自然、简洁明了（如 `A. 高筋面粉  B. 低筋面粉  C. 中筋面粉  D. 玉米淀粉`）。",
        "2. **统考经典题型呈现**：",
        "   - 题干统一采用统考标准括号填空格式：`1．（　　）题干内容……`；",
        "   - 选项短项采用制表符两列横排（`A．...  B．...`），选项长项独立分行，版面紧凑清爽。",
        "3. **第20单元完全融合教师样本题**：",
        "   - 第20单元（烘焙百分比与材料用量）直接完整收录 `unit20_quiz .docx` 中的 12 道经典选择题，并补充 3 道统考高频考查点，凑足 15 题标准卷。",
        "4. **教师卷高效批改表格**：",
        "   - 教师卷首部采用与样本一致的精炼三列批改表（`题号 ｜ 答案 ｜ 解析`），老师 10 秒内核对批改完毕，随后附带全题题干答案核对。",
        "\n---",
        "### 二、B版成果文件清单（全部存放于「厨艺理论测验题2026」）\n",
        "| 文件类型 | 数量 | 文件名称与说明 |",
        "| :--- | :---: | :--- |",
        "| **学生独立测验卷（B版）** | 20 份 | `第01单元_烘焙概论_学生测验卷_B版.docx` 至 `第20单元_烘焙百分比与材料用量_学生测验卷_B版.docx` |",
        "| **教师解析参考卷（B版）** | 20 份 | `第01单元_烘焙概论_教师解析与参考答案_B版.docx` 至 `第20单元_烘焙百分比与材料用量_教师解析与参考答案_B版.docx` |",
        "| **学生卷总汇编（B版）** | 1 份 | `高一厨艺理论_全套第01至20单元_300题学生测验卷总汇编_B版.docx`（含总目录，适合装订整册） |",
        "| **教师解析总汇编（B版）** | 1 份 | `高一厨艺理论_全套第01至20单元_300题教师详细解析总汇编_B版.docx`（含 300 题答案速查总矩阵） |",
        "| **标准JSON题库（B版）** | 1 份 | `高一厨艺理论_20单元300题标准测验题库_B版.json` |",
        "| **说明备忘录（B版）** | 1 份 | `00_高一厨艺理论测验题2026_B版_统考真题风格说明.md` |"
    ]
    with open(md_path, 'w', encoding='utf-8') as f:
        f.write("\n".join(lines))
    print(f"Generated Markdown Guide (B-Ver): {md_path}")

# -------------------------------------------------------------
# 5. MAIN EXECUTION
# -------------------------------------------------------------
def main():
    print(f">>> 1. Loaded {len(ALL_B_UNITS)} units for Version B (B版).")
    total_qs = sum(len(u['questions']) for u in ALL_B_UNITS)
    print(f">>> Total questions in B version: {total_qs}")

    print("\n>>> 2. Generating 20 Student B-version Word Documents...")
    for u in ALL_B_UNITS:
        build_student_b_doc(u, is_standalone=True)

    print("\n>>> 3. Generating 20 Teacher B-version Word Documents...")
    for u in ALL_B_UNITS:
        build_teacher_b_doc(u, is_standalone=True)

    print("\n>>> 4. Generating Master Student Compilation (B-Ver)...")
    build_master_student_b(ALL_B_UNITS)

    print("\n>>> 5. Generating Master Teacher Compilation (B-Ver)...")
    build_master_teacher_b(ALL_B_UNITS)

    print("\n>>> 6. Exporting JSON Database and Markdown Guide (B-Ver)...")
    export_b_json_and_md(ALL_B_UNITS)

    print("\n🎉 ALL VERSION B TASKS COMPLETED SUCCESSFULLY!")

if __name__ == '__main__':
    main()
