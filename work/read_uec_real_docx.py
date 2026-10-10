import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

doc_path = r'统考历届考题拆分与高一测验卷_Word\01_统考历届真题分类拆分_高一烘焙理论与实务(2020-2025).docx'
doc = docx.Document(doc_path)
print(f"Total paragraphs in UEC real questions doc: {len(doc.paragraphs)}")
for idx, p in enumerate(doc.paragraphs[:50]):
    if p.text.strip():
        print(f"[{idx}] {p.text}")
