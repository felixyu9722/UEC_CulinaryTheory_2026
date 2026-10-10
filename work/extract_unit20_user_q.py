# -*- coding: utf-8 -*-
"""
Test script to check unit20 questions from unit20_quiz .docx
"""
import docx, sys, json
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('unit20_quiz .docx')

# Extract all MCQs
mcqs = []
current_q = None

for p in doc.paragraphs:
    t = p.text.strip()
    if not t:
        continue
    # Check if starts with number like "1．（"
    if any(t.startswith(f"{i}．（") for i in range(1, 13)):
        if current_q:
            mcqs.append(current_q)
        current_q = {'q': t, 'opts': {}}
    elif current_q and '\t' in t and ('A．' in t or 'C．' in t):
        parts = t.split('\t')
        for part in parts:
            part = part.strip()
            if part.startswith('A．'): current_q['opts']['A'] = part[2:].strip()
            elif part.startswith('B．'): current_q['opts']['B'] = part[2:].strip()
            elif part.startswith('C．'): current_q['opts']['C'] = part[2:].strip()
            elif part.startswith('D．'): current_q['opts']['D'] = part[2:].strip()
    elif current_q and any(t.startswith(k + '．') for k in ['A', 'B', 'C', 'D']):
        k = t[0]
        current_q['opts'][k] = t[2:].strip()

if current_q:
    mcqs.append(current_q)

print(f"Extracted {len(mcqs)} MCQs from unit20_quiz .docx")

# Also extract answers from table 6
tbl = doc.tables[6]
ans_map = {}
for r in tbl.rows[1:]:
    q_num = int(r.cells[0].text.strip())
    ans = r.cells[1].text.strip()
    exp = r.cells[2].text.strip()
    ans_map[q_num] = {'ans': ans, 'exp': exp}

for idx, q in enumerate(mcqs):
    num = idx + 1
    a_info = ans_map.get(num, {})
    print(f"Q{num}: {q['q']}")
    print(f"   Opts: {q['opts']}")
    print(f"   Ans: {a_info.get('ans')} | Exp: {a_info.get('exp')}")
