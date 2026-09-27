from pathlib import Path
from zipfile import ZipFile
from docx import Document


def extract(path: Path, out_dir: Path) -> None:
    doc = Document(path)
    lines = [f"# SOURCE: {path.name}", ""]
    for i, p in enumerate(doc.paragraphs, 1):
        text = p.text.strip()
        if text:
            lines.append(f"[P{i:04d}][{p.style.name}] {text}")
    for ti, table in enumerate(doc.tables, 1):
        lines.append(f"\n[TABLE {ti}]")
        for ri, row in enumerate(table.rows, 1):
            cells = [c.text.replace("\n", " / ").strip() for c in row.cells]
            lines.append(f"R{ri:03d}\t" + "\t".join(cells))

    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / f"{path.stem}-extracted.txt").write_text("\n".join(lines), encoding="utf-8")

    with ZipFile(path) as zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
    (out_dir / f"{path.stem}-inventory.txt").write_text(
        f"paragraphs={len(doc.paragraphs)}\n"
        f"tables={len(doc.tables)}\n"
        f"inline_shapes={len(doc.inline_shapes)}\n"
        f"media_files={len(media)}\n" + "\n".join(media),
        encoding="utf-8",
    )


if __name__ == "__main__":
    root = Path(__file__).resolve().parent
    output = root / "work" / "docx_extract"
    for filename in ("烘焙.docx", "烘焙原料笔记总.docx"):
        extract(root / filename, output)
